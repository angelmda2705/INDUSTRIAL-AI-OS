import io

from openpyxl import load_workbook

from app.services.reports.render import safe_cell
from tests.conftest import import_csv

PROD = "Fecha,Centro de trabajo,Cantidad planeada,Piezas buenas,Scrap\n01/03/2024,L1,100,90,10\n02/03/2024,=CMD(),100,60,5\n"


def _xlsx(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_report_requires_data(admin_a):
    r = admin_a.post("/api/v1/reports", json={"report_type": "diagnostic", "file_format": "pdf"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "no_data_for_report"


def test_xlsx_report_reflects_stored_values(admin_a):
    import_csv(admin_a, PROD, "production_records", "Producción")
    admin_a.put("/api/v1/kpis/production_compliance/target", json={"target_value": 90})
    r = admin_a.post("/api/v1/reports", json={"report_type": "diagnostic", "file_format": "xlsx"})
    assert r.status_code == 201
    report = r.json()
    assert report["parameters"]["is_demo"] is False and len(report["file_sha256"]) == 64
    content = admin_a.get(f"/api/v1/reports/{report['id']}/download").content
    wb = _xlsx(content)
    assert wb.sheetnames == ["Resumen", "Indicadores", "Hallazgos", "Acciones pendientes", "Fuentes", "Limitaciones"]
    kpis = {row[0]: row for row in wb["Indicadores"].iter_rows(min_row=2, values_only=True)}
    # compliance = (90 + 60) / 200 = 75 %
    assert kpis["Cumplimiento de producción"][2] == 75.0
    assert kpis["Cumplimiento de producción"][5] == 90
    summary = {row[0]: row[1] for row in wb["Resumen"].iter_rows(values_only=True)}
    assert summary["Organización"] == "Empresa A"
    sources = list(wb["Fuentes"].iter_rows(min_row=2, values_only=True))
    assert sources[0][7] == "datos.csv"
    limitations = [r[0] for r in wb["Limitaciones"].iter_rows(min_row=2, values_only=True)]
    assert any("no son ahorros realizados" in x for x in limitations)


def test_pdf_report_and_listing(admin_a):
    import_csv(admin_a, PROD, "production_records", "Producción")
    r = admin_a.post("/api/v1/reports", json={"report_type": "executive", "file_format": "pdf"})
    assert r.status_code == 201
    download = admin_a.get(f"/api/v1/reports/{r.json()['id']}/download")
    assert download.status_code == 200 and download.content.startswith(b"%PDF")
    assert len(admin_a.get("/api/v1/reports").json()) == 1
    # Area report without matching data
    r = admin_a.post("/api/v1/reports", json={"report_type": "maintenance", "file_format": "pdf"})
    assert r.status_code == 422


def test_spreadsheet_formula_injection_is_neutralized(admin_a):
    assert safe_cell("=1+1") == "'=1+1" and safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell(5) == 5
    ds = import_csv(admin_a, PROD, "production_records", "Producción")["job"]["target_dataset_id"]
    wb = _xlsx(admin_a.get(f"/api/v1/datasets/{ds}/export").content)
    cells = [c for row in wb.active.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert "'=CMD()" in cells and "=CMD()" not in cells


def test_demo_data_is_labelled_and_removable(admin_a):
    r = admin_a.post("/api/v1/demo")
    assert r.status_code == 200, r.text
    datasets = r.json()["datasets"]
    assert len(datasets) == 4
    assert all(d["is_demo"] and d["name"].startswith("[DEMO]") for d in datasets)
    assert admin_a.post("/api/v1/demo").status_code == 409

    findings = admin_a.get("/api/v1/findings", params={"page_size": 100}).json()
    assert findings["total"] > 0 and all(f["is_demo"] for f in findings["items"])
    rules = {f["rule_code"] for f in findings["items"]}
    # The demo generator injects known situations; the engine must find them.
    assert {"FLT-FUEL-ANOMALY", "MNT-RECURRENT-FAILURES", "INV-STOCKOUT", "PRD-DOWNTIME-CONCENTRATION"} <= rules
    assert any(r.startswith("KPI-TREND:scrap_rate") for r in rules)

    dash = admin_a.get("/api/v1/dashboard").json()
    assert dash["has_demo_data"] and dash["findings"]["open_total"] == findings["total"]
    assert {k["kpi_code"] for k in dash["kpis"]} >= {"production_compliance", "mttr_hours", "fuel_efficiency_km_l"}
    assert dash["series"]
    excluded = admin_a.get("/api/v1/dashboard", params={"include_demo": False}).json()
    assert excluded["findings"]["open_total"] == 0 and excluded["kpis"] == []

    r = admin_a.post("/api/v1/reports", json={"report_type": "diagnostic", "file_format": "xlsx"})
    wb = _xlsx(admin_a.get(f"/api/v1/reports/{r.json()['id']}/download").content)
    assert wb["Resumen"]["A1"].value == "AVISO"

    inv = next(d for d in datasets if d["template_code"] == "inventory_items")
    analysis = admin_a.get(f"/api/v1/datasets/{inv['id']}/inventory-analysis").json()
    assert len(analysis["items"]) == 80 and {"A", "B", "C"} <= {i["_abc"] for i in analysis["items"]}

    assert admin_a.delete("/api/v1/demo").status_code == 200
    assert admin_a.get("/api/v1/datasets").json() == []
    assert admin_a.get("/api/v1/findings").json()["total"] == 0
    assert admin_a.get("/api/v1/imports").json()["total"] == 0
    assert admin_a.get("/api/v1/dashboard").json()["has_data"] is False


def test_dashboard_empty_state(admin_a):
    d = admin_a.get("/api/v1/dashboard").json()
    assert d["has_data"] is False and d["kpis"] == [] and d["findings"]["open_total"] == 0
    assert d["activity"]  # admins see the audit trail (login, organization creation)


def test_kpi_detail_series_and_history(admin_a):
    csv = "Fecha,Centro de trabajo,Cantidad planeada,Piezas buenas\n01/01/2024,L1,100,90\n01/02/2024,L1,100,80\n"
    ds = import_csv(admin_a, csv, "production_records", "Producción")["job"]["target_dataset_id"]
    admin_a.post(f"/api/v1/datasets/{ds}/analyze")
    d = admin_a.get("/api/v1/kpis/production_compliance", params={"dataset_id": ds}).json()
    assert d["total"]["value"] == 85.0
    assert [m["value"] for m in d["monthly"]] == [90.0, 80.0]
    assert d["by_dimension"][0]["dimension_value"] == "L1"
    assert len(d["history"]) == 2  # one value per analysis run
    assert d["definition"]["formula"]
