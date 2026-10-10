"""Render report data to PDF (ReportLab) and Excel (openpyxl)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import get_settings


def _num(v, digits=2) -> str:
    if v is None:
        return "—"
    return f"{v:,.{digits}f}"


def safe_cell(value):
    """Neutralize spreadsheet formula injection for text cells."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def render_xlsx(data: dict) -> bytes:
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="0F766E")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(title: str, header: list[str], rows: list[list]):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for r in rows:
            ws.append([safe_cell(v) for v in r])
        for i, _ in enumerate(header, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 22
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "Resumen"
    meta = [
        ("Reporte", data["title"]), ("Organización", data["organization"]), ("Periodo analizado", data["period"]),
        ("Generado", data["generated_at"] + " UTC"), ("Generado por", data["generated_by"]),
        ("Hallazgos abiertos", data["summary"]["open_findings"]),
        ("Críticos o altos", data["summary"]["critical_or_high"]),
        (f"Impacto estimado abierto ({data['currency']})", data["summary"]["estimated_impact_open"]),
        ("Indicadores calculados", data["summary"]["kpis_ok"]),
        ("Indicadores sin datos suficientes", data["summary"]["kpis_insufficient"]),
    ]
    if data["is_demo"]:
        meta.insert(0, ("AVISO", "CONTIENE DATOS DE DEMOSTRACIÓN SINTÉTICOS"))
    for k, v in meta:
        ws.append([k, safe_cell(v)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60

    sheet("Indicadores", ["Indicador", "Conjunto de datos", "Valor", "Unidad", "Estado", "Objetivo", "Evaluación",
                          "Muestra", "Fórmula", "Nota"],
          [[k["name"], k["dataset_name"], k["value"], k["unit"], "Calculado" if k["status"] == "ok" else "Datos insuficientes",
            k["target"], {"meets": "Cumple", "misses": "No cumple"}.get(k["evaluation"], "Sin objetivo"),
            k["sample_size"], k["formula"], k["message"]] for k in data["kpis"]])
    sheet("Hallazgos", ["Código", "Título", "Área", "Naturaleza", "Prioridad", "Estado", "Condición vigente", "Detectado",
                        "Descripción", "Método", "Evidencia", "Recomendación", "Impacto estimado", "Moneda", "Base del impacto",
                        "Confianza", "Responsable", "Fecha límite"],
          [[f["code"], f["title"], f["area"], f["nature"], f["priority"], f["status"], "Sí" if f["condition_active"] else "No",
            f["detected_at"], f["description"], f["method"], f["evidence_summary"], f["recommendation"], f["impact_amount"],
            f["impact_currency"], f["impact_basis"], f["confidence"], f["owner"], f["due_date"]] for f in data["findings"]])
    sheet("Acciones pendientes", ["Código", "Título", "Prioridad", "Estado", "Responsable", "Fecha límite"],
          [[f["code"], f["title"], f["priority"], f["status"], f["owner"] or "Sin asignar", f["due_date"]] for f in data["pending_actions"]])
    src_rows = []
    for s in data["sources"]:
        for file in s["files"] or [{}]:
            src_rows.append([s["dataset"], s["template"], s["rows"], "Sí" if s["is_demo"] else "No", s["period"],
                             s["analysis_run_id"], s["analyzed_at"], file.get("filename"), file.get("sha256"), file.get("rows")])
    sheet("Fuentes", ["Conjunto de datos", "Plantilla", "Registros", "Demo", "Periodo", "Análisis #", "Analizado",
                      "Archivo", "SHA-256 (inicio)", "Renglones importados"], src_rows)
    ws_lim = sheet("Limitaciones", ["Limitación"], [[x] for x in data["limitations"]])
    ws_lim.column_dimensions["A"].width = 120
    for ws_any in wb.worksheets:
        for row in ws_any.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=False, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(data: dict) -> bytes:
    settings = get_settings()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=data["title"],
                            author=settings.brand_name)
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8, leading=10, alignment=TA_LEFT)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=12)
    h1, h2 = styles["Title"], styles["Heading2"]
    warn = ParagraphStyle("warn", parent=body, textColor=colors.HexColor("#b45309"), fontName="Helvetica-Bold")

    def P(text, style=small):
        from xml.sax.saxutils import escape

        return Paragraph(escape(str(text if text is not None else "—")), style)

    def table(header, rows, widths):
        t = Table([[P(h, small) for h in header]] + [[P(c) for c in r] for r in rows], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ]))
        return t

    story = [P(settings.brand_name, small), P(data["title"], h1)]
    if data["is_demo"]:
        story.append(P("AVISO: este reporte contiene datos de demostración sintéticos que no pertenecen a la organización.", warn))
    story += [
        P(f"Organización: {data['organization']}", body),
        P(f"Periodo analizado: {data['period']}", body),
        P(f"Generado: {data['generated_at']} UTC por {data['generated_by']}", body),
        Spacer(1, 8),
        P("Resumen", h2),
        table(["Hallazgos abiertos", "Críticos/altos", f"Impacto estimado abierto ({data['currency']})",
               "Indicadores calculados", "Indicadores sin datos suficientes"],
              [[data["summary"]["open_findings"], data["summary"]["critical_or_high"], _num(data["summary"]["estimated_impact_open"]),
                data["summary"]["kpis_ok"], data["summary"]["kpis_insufficient"]]], [5 * cm] * 5),
        Spacer(1, 6),
        P("Indicadores", h2),
        table(["Indicador", "Conjunto", "Valor", "Objetivo", "Evaluación", "Muestra", "Fórmula / nota"],
              [[k["name"], k["dataset_name"], f"{_num(k['value'])} {k['unit']}" if k["status"] == "ok" else "Datos insuficientes",
                _num(k["target"]) if k["target"] is not None else "Sin objetivo",
                {"meets": "Cumple", "misses": "No cumple"}.get(k["evaluation"], "—"), k["sample_size"],
                k["formula"] if k["status"] == "ok" else k["message"]] for k in data["kpis"]],
              [5 * cm, 4.5 * cm, 3 * cm, 2.5 * cm, 2.2 * cm, 1.6 * cm, 7.8 * cm]),
        PageBreak(),
        P("Hallazgos", h2),
    ]
    if data["findings"]:
        story.append(table(
            ["Código", "Hallazgo", "Naturaleza / prioridad", "Evidencia y método", "Recomendación", "Impacto estimado", "Estado"],
            [[f["code"], f"{f['title']}. {f['description']}", f"{f['nature']} · {f['priority']}" + (f" · confianza {f['confidence']}" if f["confidence"] else ""),
              f"{f['evidence_summary']}. Método: {f['method']}", f["recommendation"],
              f"{_num(f['impact_amount'])} {f['impact_currency']} — {f['impact_basis']}" if f["impact_amount"] is not None else "No calculable",
              f["status"] + ("" if f["condition_active"] else " (ya no detectado)")] for f in data["findings"]],
            [2 * cm, 6.5 * cm, 3 * cm, 5.5 * cm, 4.5 * cm, 3.5 * cm, 1.8 * cm],
        ))
    else:
        story.append(P("No se detectaron hallazgos con las reglas vigentes.", body))
    story += [Spacer(1, 6), P("Acciones pendientes", h2)]
    if data["pending_actions"]:
        story.append(table(["Código", "Hallazgo", "Prioridad", "Estado", "Responsable", "Fecha límite"],
                           [[f["code"], f["title"], f["priority"], f["status"], f["owner"] or "Sin asignar", f["due_date"] or "—"]
                            for f in data["pending_actions"]],
                           [2.2 * cm, 11 * cm, 2.5 * cm, 3.5 * cm, 4.5 * cm, 3 * cm]))
    else:
        story.append(P("No hay acciones pendientes.", body))
    story += [Spacer(1, 6), P("Fuentes de datos", h2)]
    src_rows = []
    for s in data["sources"]:
        files = ", ".join(f"{f['filename']} (sha256 {f['sha256']}…, {f['rows']} renglones)" for f in s["files"]) or "—"
        src_rows.append([s["dataset"] + (" [DEMO]" if s["is_demo"] else ""), s["rows"], s["period"],
                         f"#{s['analysis_run_id']} · {s['analyzed_at']}" if s["analysis_run_id"] else "Sin análisis", files])
    story.append(table(["Conjunto de datos", "Registros", "Periodo", "Análisis", "Archivos importados"], src_rows,
                       [5.5 * cm, 2 * cm, 4 * cm, 4 * cm, 11 * cm]))
    story += [Spacer(1, 6), P("Limitaciones", h2)] + [P(f"• {x}", body) for x in data["limitations"]]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(1.5 * cm, 0.8 * cm, f"{data['organization']} · {data['title']} · {data['generated_at']} UTC")
        canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 0.8 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
