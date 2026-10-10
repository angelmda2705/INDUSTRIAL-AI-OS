"""Safe reading of uploaded CSV and Excel (.xlsx) files.

Values are returned *raw* (strings for CSV; native cell values for Excel) so
that type conversion happens in one place (``validation.py``) with explicit,
documented rules.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.core.errors import ValidationFailed

CSV_DELIMITERS = [",", ";", "\t", "|"]
MAX_COLUMNS = 200


@dataclass
class ParsedTable:
    columns: list[str]
    rows: list[tuple[int, list[Any]]]  # (source_row_number, values)
    sheet_names: list[str] = field(default_factory=list)
    sheet_name: str | None = None
    delimiter: str | None = None
    encoding: str | None = None

    @property
    def total_rows(self) -> int:
        return len(self.rows)


def normalize_header(text: str) -> str:
    value = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def _clean_headers(raw: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for idx, value in enumerate(raw, start=1):
        name = str(value).strip() if value is not None else ""
        name = re.sub(r"\s+", " ", name)[:120]
        if not name:
            name = f"columna_{idx}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name}_{seen[key]}"
        else:
            seen[key] = 1
        headers.append(name)
    return headers


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _decode(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValidationFailed("No se pudo leer el archivo CSV: codificación no reconocida (usa UTF-8).",
                           code="invalid_encoding")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(CSV_DELIMITERS)).delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        counts = {d: first.count(d) for d in CSV_DELIMITERS}
        best = max(counts, key=counts.get)
        return best if counts[best] > 0 else ","


def read_csv(data: bytes, max_rows: int) -> ParsedTable:
    if b"\x00" in data[:4096]:
        raise ValidationFailed("El archivo no parece ser un CSV de texto.", code="invalid_file")
    text, encoding = _decode(data)
    delimiter = _sniff_delimiter(text[:65536])
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    columns: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []
    try:
        for line_no, raw in enumerate(reader, start=1):
            if columns is None:
                if all(_is_empty(v) for v in raw):
                    continue
                if len(raw) > MAX_COLUMNS:
                    raise ValidationFailed(f"El archivo tiene más de {MAX_COLUMNS} columnas.", code="too_many_columns")
                columns = _clean_headers(raw)
                continue
            if all(_is_empty(v) for v in raw):
                continue
            values = [v.strip() for v in raw[: len(columns)]]
            values += [""] * (len(columns) - len(values))
            rows.append((line_no, values))
            if len(rows) > max_rows:
                raise ValidationFailed(
                    f"El archivo excede el máximo de {max_rows:,} renglones permitidos por importación.",
                    code="too_many_rows",
                )
    except csv.Error as exc:
        raise ValidationFailed(f"CSV mal formado: {exc}", code="invalid_csv") from exc
    if columns is None:
        raise ValidationFailed("El archivo está vacío o no tiene encabezados.", code="empty_file")
    return ParsedTable(columns=columns, rows=rows, delimiter=delimiter, encoding=encoding)


def _check_xlsx_container(data: bytes, max_uncompressed: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
                raise ValidationFailed("El archivo no es un libro de Excel .xlsx válido.", code="invalid_file")
            total = sum(info.file_size for info in zf.infolist())
            if total > max_uncompressed:
                raise ValidationFailed("El archivo de Excel es demasiado grande al descomprimirse.", code="file_too_large")
    except zipfile.BadZipFile as exc:
        raise ValidationFailed("El archivo no es un libro de Excel .xlsx válido.", code="invalid_file") from exc


def xlsx_sheet_names(data: bytes, max_uncompressed: int) -> list[str]:
    from openpyxl import load_workbook

    _check_xlsx_container(data, max_uncompressed)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_xlsx(data: bytes, max_rows: int, max_uncompressed: int, sheet: str | None = None) -> ParsedTable:
    from openpyxl import load_workbook

    _check_xlsx_container(data, max_uncompressed)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for corrupt files
        raise ValidationFailed("No se pudo abrir el libro de Excel. Verifica que no esté dañado.", code="invalid_file") from exc
    try:
        names = list(wb.sheetnames)
        if sheet is None:
            sheet = names[0]
        if sheet not in names:
            raise ValidationFailed(f"La hoja '{sheet}' no existe en el archivo.", code="sheet_not_found")
        ws = wb[sheet]
        columns: list[str] | None = None
        rows: list[tuple[int, list[Any]]] = []
        for row_no, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            raw = list(raw)
            if columns is None:
                if all(_is_empty(v) for v in raw):
                    continue
                while raw and _is_empty(raw[-1]):
                    raw.pop()
                if len(raw) > MAX_COLUMNS:
                    raise ValidationFailed(f"La hoja tiene más de {MAX_COLUMNS} columnas.", code="too_many_columns")
                columns = _clean_headers(raw)
                continue
            if all(_is_empty(v) for v in raw):
                continue
            values = [v.strip() if isinstance(v, str) else v for v in raw[: len(columns)]]
            values += [None] * (len(columns) - len(values))
            rows.append((row_no, values))
            if len(rows) > max_rows:
                raise ValidationFailed(
                    f"La hoja excede el máximo de {max_rows:,} renglones permitidos por importación.",
                    code="too_many_rows",
                )
        if columns is None:
            raise ValidationFailed("La hoja está vacía o no tiene encabezados.", code="empty_file")
        return ParsedTable(columns=columns, rows=rows, sheet_names=names, sheet_name=sheet)
    finally:
        wb.close()


def read_table(path: Path, file_type: str, *, max_rows: int, max_uncompressed: int, sheet: str | None = None) -> ParsedTable:
    data = path.read_bytes()
    if file_type == "csv":
        return read_csv(data, max_rows)
    if file_type == "xlsx":
        return read_xlsx(data, max_rows, max_uncompressed, sheet)
    raise ValidationFailed("Tipo de archivo no soportado.", code="unsupported_file")


def detect_file_type(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        if not data.startswith(b"PK\x03\x04"):
            raise ValidationFailed("El archivo tiene extensión .xlsx pero no es un libro de Excel válido.", code="invalid_file")
        return "xlsx"
    if lower.endswith(".csv") or lower.endswith(".txt"):
        return "csv"
    if lower.endswith(".xls") or lower.endswith(".xlsm"):
        raise ValidationFailed(
            "Formato no soportado. Guarda el archivo como Excel (.xlsx) sin macros o como CSV.", code="unsupported_file"
        )
    raise ValidationFailed("Formato no soportado. Usa archivos .csv o .xlsx.", code="unsupported_file")


# --- Column profiling for the preview --------------------------------------------

_NUM_RE = re.compile(r"^[\s$€]*-?[\d.,]+\s*$")
_DATE_RE = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}([ T]\d{1,2}:\d{2}(:\d{2})?)?$")


def infer_kind(values: list[Any]) -> str:
    non_empty = [v for v in values if not _is_empty(v)]
    if not non_empty:
        return "vacía"
    sample = non_empty[:300]
    n = len(sample)
    dates = sum(1 for v in sample if isinstance(v, (date, datetime)) or (isinstance(v, str) and _DATE_RE.match(v)))
    nums = sum(
        1 for v in sample
        if (isinstance(v, (int, float)) and not isinstance(v, bool)) or (isinstance(v, str) and _NUM_RE.match(v))
    )
    if dates / n >= 0.9:
        return "fecha"
    if nums / n >= 0.9:
        return "número"
    return "texto"


def column_profiles(table: ParsedTable) -> list[dict]:
    profiles = []
    for idx, name in enumerate(table.columns):
        values = [row[idx] for _, row in table.rows]
        missing = sum(1 for v in values if _is_empty(v))
        examples = []
        for v in values:
            if not _is_empty(v):
                text = v.isoformat() if isinstance(v, (date, datetime)) else str(v)
                if text not in examples:
                    examples.append(text[:60])
            if len(examples) >= 3:
                break
        profiles.append(
            {"name": name, "kind": infer_kind(values), "missing": missing, "examples": examples}
        )
    return profiles


def jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value != value:  # NaN
        return None
    return value
